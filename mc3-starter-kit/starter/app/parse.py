"""Walk a corpus once and write one JSON chunk per fact.

One bad file must not stop the walk. The graded corpus contains an
unreadable file, an encrypted PDF, an unknown type, and an empty directory.
"""

from __future__ import annotations

import csv
import json
import os
import re
import subprocess
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".log"}
ASSIGN_RE = re.compile(r"^[A-Z_][A-Z0-9_]*\s*=")
READ_CAP = 5_000_000


def build_index(corpus: Path, index_dir: Path) -> list[dict]:
    """Parse corpus into index_dir/chunks.jsonl. Return the chunks."""
    corpus = corpus.resolve()
    index_dir.mkdir(parents=True, exist_ok=True)
    chunks: list[dict] = []
    skipped: list[dict] = []

    def note(path: Path, reason: str) -> None:
        rel = _rel(corpus, path)
        skipped.append({"path": rel, "reason": reason})

    for dirpath, _dirs, filenames in os.walk(corpus, onerror=lambda err: note(Path(err.filename), type(err).__name__)):
        for name in filenames:
            path = Path(dirpath) / name
            try:
                chunks.extend(_parse_file(corpus, path))
            except (KeyboardInterrupt, SystemExit):
                raise
            except Exception as exc:
                note(path, type(exc).__name__)

    _write_jsonl(index_dir / "chunks.jsonl", chunks)
    (index_dir / "skipped.json").write_text(
        json.dumps(skipped, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return chunks


def load_chunks(index_dir: Path) -> list[dict]:
    path = index_dir / "chunks.jsonl"
    if not path.exists():
        return []
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line:
            rows.append(json.loads(line))
    return rows


def save_chunks(index_dir: Path, chunks: list[dict]) -> None:
    _write_jsonl(index_dir / "chunks.jsonl", chunks)


def _parse_file(corpus: Path, path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    rel = _rel(corpus, path)
    if suffix == ".pdf":
        return _pdf(path, rel)
    if suffix == ".docx":
        return _docx(path, rel)
    if suffix == ".xlsx":
        return _xlsx(path, rel)
    if suffix == ".csv":
        return _csv(path, rel)
    if suffix in TEXT_SUFFIXES:
        return _lines(path, rel, "line")
    if suffix == ".py":
        return _python(path, rel)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return _image(path, rel)
    raise RuntimeError("unknown-type")


def _pdf(path: Path, rel: str) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted and reader.decrypt("") == 0:
        raise RuntimeError("encrypted")
    pages = [(page.extract_text() or "") for page in reader.pages]
    if not any(page.strip() for page in pages):
        raise RuntimeError("encrypted")
    flag = _withdrawn(path, pages[0] if pages else "")
    return [_chunk(rel, page, "pdf", flag) for page in pages if page.strip()]


def _docx(path: Path, rel: str) -> list[dict]:
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            # A merged cell repeats the same text. Keep one copy of each.
            seen: list[str] = []
            for cell in cells:
                if cell and cell not in seen:
                    seen.append(cell)
            if seen:
                rows.append(" | ".join(seen))
    flag = _withdrawn(path, "\n".join(paragraphs[:3]))
    chunks = [_chunk(rel, text, "paragraph", flag) for text in paragraphs]
    chunks.extend(_chunk(rel, text, "table", flag) for text in rows)
    return chunks


def _xlsx(path: Path, rel: str) -> list[dict]:
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        chunks: list[dict] = []
        first_text = ""
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                vals = ["" if value is None else str(value).strip() for value in row]
                if any(vals):
                    rows.append(vals)
            if not rows:
                continue
            if not first_text:
                first_text = " ".join(rows[0])
            if len(rows) == 1:
                chunks.append(_chunk(rel, f"sheet: {sheet.title} | " + " | ".join(v for v in rows[0] if v), "row", False))
                continue
            header = rows[0]
            for row in rows[1:]:
                parts = [f"sheet: {sheet.title}"]
                for key, value in zip(header, row):
                    if not value:
                        continue
                    parts.append(f"{key}: {value}" if key else value)
                chunks.append(_chunk(rel, " | ".join(parts), "row", False))
        flag = _withdrawn(path, first_text)
        if flag:
            for chunk in chunks:
                chunk["withdrawn"] = True
        return chunks
    finally:
        book.close()


def _csv(path: Path, rel: str) -> list[dict]:
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        rows = list(csv.DictReader(handle))
    chunks = []
    for row in rows:
        parts = [f"{key}: {value}" for key, value in row.items() if value]
        if parts:
            chunks.append(_chunk(rel, " | ".join(parts), "row", False))
    return chunks


def _lines(path: Path, rel: str, kind: str) -> list[dict]:
    text = _read(path)
    flag = _withdrawn(path, text)
    return [_chunk(rel, line.strip(), kind, flag) for line in text.splitlines() if line.strip()]


def _python(path: Path, rel: str) -> list[dict]:
    text = _read(path)
    flag = _withdrawn(path, text)
    chunks = [_chunk(rel, text, "code", flag)]
    for line in text.splitlines():
        if ASSIGN_RE.match(line.strip()):
            chunks.append(_chunk(rel, line.strip(), "assign", flag))
    return chunks


def _image(path: Path, rel: str) -> list[dict]:
    transcript = _tesseract(path)
    chunk = _chunk(rel, transcript, "image", _withdrawn(path, transcript))
    chunk["needs_vision"] = True
    return [chunk]


def _tesseract(path: Path) -> str:
    try:
        proc = subprocess.run(
            ["tesseract", str(path), "stdout"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return ""
    return proc.stdout.strip()


def _withdrawn(path: Path, text: str) -> bool:
    """The file itself is withdrawn, not a later file that mentions an old one.

    Revision 2 says "Revision 1 is withdrawn" in the body. Only the file
    name and the first three non-empty lines count.
    """
    if "WITHDRAWN" in path.name.upper():
        return True
    lines = [line.strip() for line in text.splitlines() if line.strip()][:3]
    head = "\n".join(lines).upper()
    return "WITHDRAWN" in head or "SUPERSEDED" in head


def _chunk(path: str, text: str, kind: str, withdrawn: bool) -> dict:
    return {
        "path": path,
        "text": text,
        "kind": kind,
        "withdrawn": withdrawn,
        "needs_vision": False,
    }


def _read(path: Path) -> str:
    data = path.read_bytes()[:READ_CAP]
    return data.decode("utf-8", errors="replace")


def _rel(corpus: Path, path: Path) -> str:
    try:
        return path.resolve().relative_to(corpus).as_posix()
    except ValueError:
        return path.name


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    payload = "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(payload, encoding="utf-8")
    os.replace(tmp, path)
